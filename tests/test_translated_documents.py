"""Emails, SMS, PDFs, spreadsheets and CSV follow the reader's language."""
import csv
import io
from datetime import datetime

import openpyxl

from backend.i18n import Code, Phrase, render_message, use_language
from backend.models import Notification, RentPayment
from backend.scheduler_jobs import job_lease_expiry_alerts
from backend.services.notifications import send_email, send_sms
from tests.conftest import login


def _payment(db, lease, amount=100.0):
    payment = RentPayment(
        lease_id=lease.id, amount=amount, method="BANK_TRANSFER", paid_at=datetime.utcnow(), receipt_number="RCP-TEST-1",
        recorded_by=lease.unit.property.owner_id,
    )
    db.session.add(payment)
    db.session.commit()
    return payment


# --- Saved language ------------------------------------------------------------

def test_switching_language_while_signed_in_saves_it(client, db, owner):
    login(client, owner.email)
    client.get("/lang/ar")
    db.session.refresh(owner)
    assert owner.language == "ar"


def test_saved_language_follows_the_user_to_a_new_browser(app, db, owner):
    owner.language = "ar"
    db.session.commit()
    fresh_browser = app.test_client()
    page = login(fresh_browser, owner.email).get_data(as_text=True)
    assert 'dir="rtl"' in page


def test_language_picked_before_signing_in_becomes_the_saved_one(client, db, owner):
    client.get("/lang/sn")
    login(client, owner.email)
    db.session.refresh(owner)
    assert owner.language == "sn"


# --- Emails and SMS -------------------------------------------------------------

def test_email_is_written_in_the_recipients_language(app, db, owner, tenant):
    tenant.user.language = "ar"
    db.session.commit()
    arabic = send_email(tenant.user, "Payment received", "{name} paid {amount} online.", name="Tara", amount="50.00")
    english = send_email(owner, "Payment received", "{name} paid {amount} online.", name="Tara", amount="50.00")
    assert arabic.subject == "تم استلام الدفعة"
    assert arabic.body == "دفع Tara مبلغ 50.00 إلكترونيًا."
    assert english.subject == "Payment received"
    assert english.body == "Tara paid 50.00 online."


def test_sms_is_written_in_the_recipients_language(app, db, tenant):
    tenant.user.language = "ar"
    db.session.commit()
    sms = send_sms(tenant.user, "Ticket {ticket} has been received.", ticket="MNT-1")
    assert sms.body == "تم استلام التذكرة MNT-1."


def test_nightly_job_writes_to_each_person_in_their_own_language(app, db, active_lease):
    from datetime import date, timedelta

    active_lease.end_date = date.today() + timedelta(days=30)
    active_lease.tenant.user.language = "ar"
    db.session.commit()
    assert job_lease_expiry_alerts() == 1
    to_tenant = Notification.query.filter_by(recipient_id=active_lease.tenant.user_id).one()
    to_owner = Notification.query.filter_by(recipient_id=active_lease.unit.property.owner_id).one()
    assert to_tenant.subject == "عقد إيجارك على وشك الانتهاء"
    assert to_owner.subject == "Tenant lease expiring in 30 days"


def test_codes_and_phrases_inside_messages_are_translated_too(app):
    with app.test_request_context(), use_language("ar"):
        status = render_message("Ticket {ticket} is now {status}.", ticket="T1", status=Code("IN_PROGRESS"))
        tier = render_message(
            "[{tier}] Rent of {amount} is {days} day(s) overdue.", tier=Phrase("Gentle reminder"), amount="9.00", days=2
        )
    assert status == "حالة التذكرة T1 الآن: قيد التنفيذ."
    assert tier == "[تذكير ودّي] الإيجار البالغ 9.00 متأخر منذ 2 يوم/أيام."
    with app.test_request_context():
        assert render_message("Ticket {ticket} is now {status}.", ticket="T1", status=Code("IN_PROGRESS")) == (
            "Ticket T1 is now In Progress."
        )


# --- PDFs -------------------------------------------------------------------------

def test_arabic_pdfs_embed_an_arabic_font(client, db, owner, active_lease):
    payment = _payment(db, active_lease)
    login(client, owner.email)
    english = client.get(f"/rent/payments/{payment.id}/receipt.pdf").data
    client.get("/lang/ar")
    arabic = client.get(f"/rent/payments/{payment.id}/receipt.pdf").data
    assert arabic.startswith(b"%PDF") and b"DejaVuSans" in arabic
    assert b"DejaVuSans" not in english and b"Helvetica" in english


def test_every_pdf_renders_in_arabic(client, db, owner, active_lease, property_):
    _payment(db, active_lease)
    login(client, owner.email)
    client.get("/lang/ar")
    for url in (
        f"/rent/{active_lease.id}/statement.pdf",
        f"/rent/{active_lease.id}/statement.pdf?start=2026-01-01&end=2026-12-31",
        "/reports/export.pdf",
        f"/reports/export.pdf?property_id={property_.id}",
    ):
        resp = client.get(url)
        assert resp.status_code == 200 and resp.data.startswith(b"%PDF"), url


def test_names_with_markup_characters_do_not_break_pdfs(client, db, owner, active_lease):
    active_lease.tenant.user.first_name = "Tom & <Jerry>"
    db.session.commit()
    payment = _payment(db, active_lease)
    login(client, owner.email)
    for lang in ("en", "ar"):
        client.get(f"/lang/{lang}")
        assert client.get(f"/rent/payments/{payment.id}/receipt.pdf").status_code == 200


def test_arabic_pdf_text_is_joined_reordered_and_keeps_brackets_and_dates_right():
    from backend.services.pdf import _Writer

    # Display order (left to right), as drawn on the page.
    assert _Writer._visual("جميع العقارات (المحفظة)") == "(ﺔﻈﻔﺤﻤﻟﺍ) ﺕﺍﺭﺎﻘﻌﻟﺍ ﻊﻴﻤﺟ"
    assert _Writer._visual("الوحدة: \u2066U1 (Eastlea Flats)\u2069") == "U1 (Eastlea Flats) :ﺓﺪﺣﻮﻟﺍ"
    assert _Writer._visual("من \u20662026-05-02\u2069 إلى \u20662027-05-02\u2069") == "2027-05-02 ﻰﻟﺇ 2026-05-02 ﻦﻣ"


def test_property_report_pdf_shows_metric_names_not_keys(app, property_):
    from backend.services.pdf import render_property_report_pdf
    from backend.services.reports import property_performance

    with app.test_request_context():
        pdf = render_property_report_pdf(property_, property_performance(property_))
    assert b"total_units" not in pdf


# --- Spreadsheets and CSV ----------------------------------------------------------

def test_arabic_excel_report_is_translated_and_right_to_left(client, owner, property_):
    login(client, owner.email)
    client.get("/lang/ar")
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/reports/export.xlsx").data))
    assert wb.sheetnames == ["ملخص المحفظة", "حسب العقار"]
    sheet = wb["حسب العقار"]
    assert sheet.sheet_view.rightToLeft
    assert sheet["A1"].value == "العقار"


def test_english_excel_report_is_unchanged(client, owner, property_):
    login(client, owner.email)
    wb = openpyxl.load_workbook(io.BytesIO(client.get("/reports/export.xlsx").data))
    assert wb.sheetnames == ["Portfolio Summary", "By Property"]
    assert not wb["By Property"].sheet_view.rightToLeft


def test_csv_export_has_translated_headings_excel_can_read(client, db, owner, active_lease):
    _payment(db, active_lease)
    login(client, owner.email)
    client.get("/lang/ar")
    text = client.get(f"/rent/{active_lease.id}/history.csv").get_data(as_text=True)
    assert text.startswith("﻿")
    header, row = list(csv.reader(io.StringIO(text.lstrip("﻿"))))[:2]
    assert header[0] == "التاريخ" and header[4] == "رقم الإيصال"
    assert row[3] == "تحويل بنكي"
