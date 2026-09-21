"""Excel (.xlsx) report export via openpyxl (FR-043 / UC-25)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_header(ws, headers):
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def render_portfolio_report_xlsx(summary: dict) -> bytes:
    wb = Workbook()

    overview = wb.active
    overview.title = "Portfolio Summary"
    _write_header(overview, ["Metric", "Value"])
    for key in ("total_units", "occupied_units", "occupancy_rate", "rent_collected", "rent_outstanding", "maintenance_cost"):
        if key in summary:
            overview.append([key.replace("_", " ").title(), summary[key]])
    overview.column_dimensions["A"].width = 24
    overview.column_dimensions["B"].width = 18

    by_property = wb.create_sheet("By Property")
    _write_header(
        by_property,
        ["Property", "Code", "Total Units", "Occupied", "Occupancy %", "Rent Collected", "Outstanding", "Maintenance Cost", "Open Requests"],
    )
    for p in summary.get("properties", []):
        by_property.append([
            p["property_name"], p["property_code"], p["total_units"], p["occupied_units"],
            p["occupancy_rate"], p["rent_collected"], p["rent_outstanding"], p["maintenance_cost"],
            p["open_maintenance_requests"],
        ])
    for col, width in zip("ABCDEFGHI", (24, 12, 10, 10, 12, 14, 12, 16, 14)):
        by_property.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
